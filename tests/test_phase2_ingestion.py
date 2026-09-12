"""
test_phase2_ingestion.py

Tests for document ingestion (PDF/DOCX parsing) and SQLite metadata/page
storage. Uses real temporary PDF and DOCX files generated on the fly with
PyMuPDF and python-docx so tests exercise the actual parsing libraries
rather than mocks.
"""

from __future__ import annotations

import sqlite3
import uuid
from pathlib import Path

import docx
import fitz
import pytest

from app.database import repository
from app.database.exceptions import DocumentNotFoundError, DuplicateDocumentError
from app.database.sqlite_db import get_connection, init_db
from app.ingestion.document_parser import ingest_document
from app.ingestion.exceptions import (
    EmptyDocumentError,
    UnsupportedFileTypeError,
)
from app.ingestion.file_storage import save_uploaded_file
from app.models.document import DocumentType


# ----------------------------------------------------------------------------
# Fixtures: build real PDF/DOCX files in a temp directory for each test.
# ----------------------------------------------------------------------------


@pytest.fixture()
def sample_pdf(tmp_path: Path) -> Path:
    """Create a small 2-page PDF with distinct text per page."""
    pdf_path = tmp_path / "sample_contract.pdf"
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((72, 72), "PAGE ONE: Payment shall be due within 30 days of invoice.")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "PAGE TWO: Either party may terminate with 60 days written notice.")
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path


@pytest.fixture()
def sample_docx_single_page(tmp_path: Path) -> Path:
    """Create a DOCX with no explicit page breaks (single pseudo-page)."""
    docx_path = tmp_path / "sample_contract.docx"
    document = docx.Document()
    document.add_paragraph("This Agreement is made between Acme Corp and Globex Inc.")
    document.add_paragraph("Payment shall be due within 30 days of invoice.")
    document.save(str(docx_path))
    return docx_path


@pytest.fixture()
def sample_docx_multi_page(tmp_path: Path) -> Path:
    """Create a DOCX with an explicit hard page break between two sections."""
    docx_path = tmp_path / "sample_multi_page.docx"
    document = docx.Document()
    document.add_paragraph("PAGE ONE: Confidentiality obligations apply to both parties.")
    document.add_page_break()
    document.add_paragraph("PAGE TWO: Liability is limited to direct damages only.")
    document.save(str(docx_path))
    return docx_path


@pytest.fixture()
def empty_pdf(tmp_path: Path) -> Path:
    """Create a PDF with a blank page (no extractable text)."""
    pdf_path = tmp_path / "empty.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path


@pytest.fixture()
def unsupported_file(tmp_path: Path) -> Path:
    """Create a .txt file, which is not in the allowed extensions list."""
    txt_path = tmp_path / "notes.txt"
    txt_path.write_text("This is not a contract format we support.")
    return txt_path


@pytest.fixture()
def isolated_db(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    """
    Point the database layer at a temporary SQLite file for the duration of
    the test, so tests never touch the real data/contracts.db, and clear the
    cached settings so the new path takes effect.
    """
    from app.utils import config_loader

    config_loader.get_settings.cache_clear()
    settings = config_loader.get_settings()
    test_db_rel_path = f"data/test_{uuid.uuid4().hex}.db"
    settings["paths"]["sqlite_db_path"] = test_db_rel_path

    init_db()
    yield
    config_loader.get_settings.cache_clear()


# ----------------------------------------------------------------------------
# Ingestion tests
# ----------------------------------------------------------------------------


class TestPDFIngestion:
    def test_ingest_pdf_extracts_pages_and_metadata(self, sample_pdf: Path) -> None:
        parsed = ingest_document(sample_pdf, contract_name="Sample PDF Contract")

        assert parsed.metadata.contract_name == "Sample PDF Contract"
        assert parsed.metadata.document_type == DocumentType.PDF
        assert parsed.metadata.page_count == 2
        assert len(parsed.pages) == 2
        assert "PAGE ONE" in parsed.pages[0].text
        assert "PAGE TWO" in parsed.pages[1].text
        assert parsed.pages[0].page_number == 1
        assert parsed.pages[1].page_number == 2

    def test_ingest_pdf_defaults_contract_name_to_filename(self, sample_pdf: Path) -> None:
        parsed = ingest_document(sample_pdf)
        assert parsed.metadata.contract_name == "sample_contract"

    def test_empty_pdf_raises(self, empty_pdf: Path) -> None:
        with pytest.raises(EmptyDocumentError):
            ingest_document(empty_pdf)


class TestDOCXIngestion:
    def test_ingest_docx_single_page(self, sample_docx_single_page: Path) -> None:
        parsed = ingest_document(sample_docx_single_page, contract_name="Sample DOCX Contract")

        assert parsed.metadata.document_type == DocumentType.DOCX
        assert parsed.metadata.page_count == 1
        assert "Acme Corp" in parsed.pages[0].text
        assert "Payment shall be due" in parsed.pages[0].text

    def test_ingest_docx_multi_page_via_page_breaks(self, sample_docx_multi_page: Path) -> None:
        parsed = ingest_document(sample_docx_multi_page, contract_name="Multi Page DOCX")

        assert parsed.metadata.page_count == 2
        assert "PAGE ONE" in parsed.pages[0].text
        assert "Confidentiality" in parsed.pages[0].text
        assert "PAGE TWO" in parsed.pages[1].text
        assert "Liability is limited" in parsed.pages[1].text


class TestValidation:
    def test_unsupported_extension_rejected(self, unsupported_file: Path) -> None:
        with pytest.raises(UnsupportedFileTypeError):
            ingest_document(unsupported_file)


class TestFileStorage:
    def test_save_uploaded_file_writes_bytes_to_disk(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.utils import config_loader

        config_loader.get_settings.cache_clear()
        settings = config_loader.get_settings()
        # Use an absolute path override; save_uploaded_file joins project_root / raw_documents_dir,
        # and Path(project_root) / absolute_path resolves to the absolute path itself.
        settings["paths"]["raw_documents_dir"] = str(tmp_path)

        saved_path = save_uploaded_file(b"%PDF-1.4 fake content", "my contract.pdf")
        assert saved_path.exists()
        assert saved_path.read_bytes() == b"%PDF-1.4 fake content"
        assert saved_path.name.endswith("my contract.pdf")

        config_loader.get_settings.cache_clear()


# ----------------------------------------------------------------------------
# Database / repository tests
# ----------------------------------------------------------------------------


class TestRepository:
    def test_insert_and_retrieve_round_trip(self, isolated_db, sample_pdf: Path) -> None:
        parsed = ingest_document(sample_pdf, contract_name="Round Trip Contract")
        repository.insert_parsed_document(parsed)

        retrieved = repository.get_parsed_document(parsed.metadata.document_id)
        assert retrieved.metadata.contract_name == "Round Trip Contract"
        assert retrieved.metadata.page_count == 2
        assert len(retrieved.pages) == 2
        assert retrieved.pages[0].text == parsed.pages[0].text

    def test_duplicate_insert_raises(self, isolated_db, sample_pdf: Path) -> None:
        parsed = ingest_document(sample_pdf, contract_name="Dup Contract")
        repository.insert_parsed_document(parsed)
        with pytest.raises(DuplicateDocumentError):
            repository.insert_parsed_document(parsed)

    def test_get_nonexistent_document_raises(self, isolated_db) -> None:
        with pytest.raises(DocumentNotFoundError):
            repository.get_document_metadata("does-not-exist")

    def test_list_documents_returns_all(self, isolated_db, sample_pdf: Path, sample_docx_single_page: Path) -> None:
        parsed_pdf = ingest_document(sample_pdf, contract_name="PDF Contract")
        parsed_docx = ingest_document(sample_docx_single_page, contract_name="DOCX Contract")
        repository.insert_parsed_document(parsed_pdf)
        repository.insert_parsed_document(parsed_docx)

        all_docs = repository.list_documents()
        names = {d.contract_name for d in all_docs}
        assert names == {"PDF Contract", "DOCX Contract"}

    def test_document_exists(self, isolated_db, sample_pdf: Path) -> None:
        parsed = ingest_document(sample_pdf, contract_name="Existence Check")
        assert repository.document_exists(parsed.metadata.document_id) is False
        repository.insert_parsed_document(parsed)
        assert repository.document_exists(parsed.metadata.document_id) is True

    def test_delete_document_cascades_pages(self, isolated_db, sample_pdf: Path) -> None:
        parsed = ingest_document(sample_pdf, contract_name="To Be Deleted")
        repository.insert_parsed_document(parsed)
        document_id = parsed.metadata.document_id

        repository.delete_document(document_id)

        assert repository.document_exists(document_id) is False
        with pytest.raises(DocumentNotFoundError):
            repository.get_document_metadata(document_id)

        # Verify pages were cascade-deleted, not orphaned.
        with get_connection() as conn:
            row = conn.execute(
                "SELECT COUNT(*) as cnt FROM document_pages WHERE document_id = ?", (document_id,)
            ).fetchone()
        assert row["cnt"] == 0

    def test_delete_nonexistent_document_raises(self, isolated_db) -> None:
        with pytest.raises(DocumentNotFoundError):
            repository.delete_document("ghost-id")

    def test_schema_rejects_invalid_document_type(self, isolated_db) -> None:
        with get_connection() as conn:
            with pytest.raises(sqlite3.IntegrityError):
                conn.execute(
                    """
                    INSERT INTO documents
                        (document_id, contract_name, document_type, upload_date,
                         page_count, file_size_bytes, source_path)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    ("bad-doc", "Bad Doc", "xlsx", "2026-01-01T00:00:00", 1, 100, "x.xlsx"),
                )
